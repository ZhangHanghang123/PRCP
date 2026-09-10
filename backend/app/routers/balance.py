"""资产负债表 API — 单大表 + 24 月现金流缺口"""
import io
import re
from typing import Optional, List
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy import text
from sqlalchemy.orm import Session
from pydantic import BaseModel
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment

from app.database import get_db
from app.auth import get_current_user

router = APIRouter(prefix="/balance", tags=["资产负债表"])

GAP_COLS = [f"m{i}_gap" for i in range(1, 25)]
MEASURE_COLS = ["begin_balance", "avg_balance", "interest_rate",
               "interest_amount", "capital_ratio", "risk_weight"]

# 导出/导入字段顺序（与前端 MEASURES 一致）
EXPORT_MEASURES = [
    {"key": "begin_balance",   "name": "月初余额"},
    {"key": "current_amount",  "name": "月末余额"},
    {"key": "avg_balance",     "name": "平均余额"},
    {"key": "interest_rate",   "name": "利率(%)"},
    {"key": "interest_amount", "name": "利息收支"},
    {"key": "capital_ratio",   "name": "资本占用(%)"},
    {"key": "risk_weight",     "name": "风险权重(%)"},
]


class BalanceIn(BaseModel):
    coa_node_id: int
    data_date: str  # YYYY-MM-DD
    current_amount: float = 0
    begin_balance: Optional[float] = 0
    avg_balance: Optional[float] = 0
    interest_rate: Optional[float] = 0
    interest_amount: Optional[float] = 0
    capital_ratio: Optional[float] = 0
    risk_weight: Optional[float] = 0
    gaps: Optional[List[float]] = None  # 24 个数
    calc_note: Optional[str] = None


@router.get("/")
async def list_balance(
    coa_node_id: Optional[int] = None,
    data_date: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    scheme_id: Optional[int] = None,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    where = ["b.is_deleted=0"]
    params = {}
    if coa_node_id:
        where.append("b.coa_node_id=:n")
        params["n"] = coa_node_id
    if data_date:
        where.append("b.data_date=:d")
        params["d"] = data_date
    if start_date:
        where.append("b.data_date>=:sd")
        params["sd"] = start_date
    if end_date:
        where.append("b.data_date<=:ed")
        params["ed"] = end_date
    if scheme_id:
        where.append("n.scheme_id=:s")
        params["s"] = scheme_id

    rows = db.execute(
        text(f"""SELECT b.id, b.coa_node_id, n.node_code, n.node_name, n.node_level,
                       n.path, n.node_type,
                       b.data_date, b.current_amount,
                       b.begin_balance, b.avg_balance, b.interest_rate,
                       b.interest_amount, b.capital_ratio, b.risk_weight,
                       {','.join('b.' + c for c in GAP_COLS)},
                       b.calc_note, b.created_at
                FROM prcp_data_balance b
                LEFT JOIN prcp_coa_node n ON n.id=b.coa_node_id
                WHERE {' AND '.join(where)}
                ORDER BY n.path, b.data_date DESC, b.coa_node_id
                LIMIT 2000"""),
        params,
    ).fetchall()

    items = []
    for r in rows:
        gaps = [float(r[15 + i]) for i in range(24)]
        items.append({
            "id": r[0], "coa_node_id": r[1],
            "node_code": r[2], "node_name": r[3],
            "node_level": r[4], "path": r[5], "node_type": r[6],
            "data_date": r[7].isoformat() if r[7] else None,
            "current_amount": float(r[8] or 0),
            "begin_balance": float(r[9] or 0),
            "avg_balance": float(r[10] or 0),
            "interest_rate": float(r[11] or 0),
            "interest_amount": float(r[12] or 0),
            "capital_ratio": float(r[13] or 0),
            "risk_weight": float(r[14] or 0),
            "gaps": gaps,
            "calc_note": r[39], "created_at": r[40].isoformat() if r[40] else None,
        })
    return {"items": items, "total": len(items)}


@router.post("/")
async def upsert_balance(p: BalanceIn, db: Session = Depends(get_db), user=Depends(get_current_user)):
    """按 (coa_node_id, data_date) 唯一，存在则更新，否则插入"""
    uid = user.get("id", 1) if isinstance(user, dict) else getattr(user, "id", 1)
    gaps = p.gaps or [0] * 24
    while len(gaps) < 24:
        gaps.append(0)
    gap_values = gaps[:24]

    gap_set = ",".join(f"{c}=:{c}" for c in GAP_COLS)
    gap_insert_cols = ",".join(GAP_COLS)
    gap_insert_vals = ",".join(f":{c}" for c in GAP_COLS)
    measure_set = ",".join(f"{c}=:{c}" for c in MEASURE_COLS)
    measure_insert_cols = ",".join(MEASURE_COLS)
    measure_insert_vals = ",".join(f":{c}" for c in MEASURE_COLS)

    existing = db.execute(
        text("SELECT id FROM prcp_data_balance WHERE coa_node_id=:n AND data_date=:d AND is_deleted=0"),
        {"n": p.coa_node_id, "d": p.data_date},
    ).first()
    if existing:
        params = {c: gap_values[i] for i, c in enumerate(GAP_COLS)}
        params.update({c: getattr(p, c) or 0 for c in MEASURE_COLS})
        params.update({"amt": p.current_amount, "note": p.calc_note, "u": uid, "id": existing[0]})
        db.execute(
            text(f"""UPDATE prcp_data_balance SET
                current_amount=:amt, {measure_set}, {gap_set}, calc_note=:note, updated_by=:u
                WHERE id=:id"""),
            params,
        )
        return {"id": existing[0], "action": "updated"}
    else:
        params = {c: gap_values[i] for i, c in enumerate(GAP_COLS)}
        params.update({c: getattr(p, c) or 0 for c in MEASURE_COLS})
        params.update({
            "n": p.coa_node_id, "d": p.data_date,
            "amt": p.current_amount, "note": p.calc_note, "u": uid,
        })
        rid = db.execute(
            text(f"""INSERT INTO prcp_data_balance
                (coa_node_id, data_date, current_amount,
                 {measure_insert_cols}, {gap_insert_cols},
                 calc_note, created_by, updated_by)
                VALUES (:n, :d, :amt, {measure_insert_vals}, {gap_insert_vals}, :note, :u, :u)"""),
            params,
        ).lastrowid
        return {"id": rid, "action": "created"}


@router.delete("/{bid}")
async def delete_balance(bid: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    uid = user.get("id", 1) if isinstance(user, dict) else getattr(user, "id", 1)
    r = db.execute(
        text("UPDATE prcp_data_balance SET is_deleted=1, updated_by=:u WHERE id=:id AND is_deleted=0"),
        {"u": uid, "id": bid},
    ).rowcount
    if r == 0:
        raise HTTPException(404, "记录不存在")
    return {"ok": True}


@router.get("/gap-summary")
async def gap_summary(
    data_date: str = Query(..., description="YYYY-MM-DD"),
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    """按日期聚合：每个节点 + 24 月缺口合计"""
    rows = db.execute(
        text(f"""SELECT b.coa_node_id, n.node_code, n.node_name,
                       b.current_amount,
                       {','.join('b.' + c for c in GAP_COLS)}
                FROM prcp_data_balance b
                LEFT JOIN prcp_coa_node n ON n.id=b.coa_node_id
                WHERE b.data_date=:d AND b.is_deleted=0
                ORDER BY b.coa_node_id"""),
        {"d": data_date},
    ).fetchall()
    items = []
    for r in rows:
        gaps = [float(r[4 + i]) for i in range(24)]
        items.append({
            "coa_node_id": r[0], "node_code": r[1], "node_name": r[2],
            "current_amount": float(r[3]), "gaps": gaps,
            "sum_24m": sum(gaps), "min_gap": min(gaps), "max_gap": max(gaps),
        })
    return {"data_date": data_date, "items": items, "total": len(items)}


@router.get("/by-scheme-matrix")
async def by_scheme_matrix(
    scheme_id: int = Query(...),
    start_date: str = Query(..., description="YYYY-MM-DD 起始月"),
    end_date: str = Query(..., description="YYYY-MM-DD 结束月"),
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    """二维矩阵：行 = 账户册 × 列 = 月份，单元格 = 7 个度量

    返回结构：
    - dates: ['2026-01', '2026-02', ..., '2026-12']  (按月排序)
    - nodes: [{coa_node_id, node_code, node_name, path, node_level, category}]
    - matrix: {coa_node_id: {date_str: {begin_balance, avg_balance, current_amount,
                                       interest_rate, interest_amount,
                                       capital_ratio, risk_weight}}}
    - categories: {category: {date_str: 聚合后的当月汇总值}}
    """
    # 1) 加载该方案下所有节点（按层级排序，方便按大类聚合）
    node_rows = db.execute(
        text("""SELECT id, node_code, node_name, parent_id, node_level, node_type, path, sort_order, description
                FROM prcp_coa_node
                WHERE scheme_id=:s AND is_deleted=0
                ORDER BY sort_order, path"""),
        {"s": scheme_id},
    ).fetchall()

    nodes = []
    for r in node_rows:
        path = r[6] or ""
        cat = ""
        if "/" in path:
            cat = path.split("/")[1].replace("L1_", "")
        nodes.append({
            "coa_node_id": r[0], "node_code": r[1], "node_name": r[2],
            "parent_id": r[3], "node_level": r[4], "node_type": r[5], "path": r[6],
            "sort_order": r[7] or 0,
            "category": cat,
            "description": r[8] or "",
        })

    # 计算月份列表（按月递增）
    from datetime import datetime
    from dateutil.relativedelta import relativedelta
    start = datetime.strptime(start_date, "%Y-%m-%d").replace(day=1)
    end = datetime.strptime(end_date, "%Y-%m-%d").replace(day=1)
    dates = []
    cur = start
    while cur <= end:
        dates.append(cur.strftime("%Y-%m"))
        cur = cur + relativedelta(months=1)
    if not dates:
        return {"dates": [], "nodes": [], "matrix": {}, "categories": {}}

    # 2) 一次性查所有月份的 balance（用 BETWEEN）
    rows = db.execute(
        text(f"""SELECT b.coa_node_id, DATE_FORMAT(b.data_date, '%Y-%m') AS ym,
                       b.begin_balance, b.avg_balance, b.current_amount,
                       b.interest_rate, b.interest_amount,
                       b.capital_ratio, b.risk_weight
                FROM prcp_data_balance b
                JOIN prcp_coa_node n ON n.id=b.coa_node_id
                WHERE n.scheme_id=:s AND b.is_deleted=0
                  AND b.data_date BETWEEN :sd AND :ed
                ORDER BY b.coa_node_id, b.data_date"""),
        {"s": scheme_id, "sd": start_date, "ed": end_date},
    ).fetchall()

    # 3) 构建 matrix[coa_node_id][ym] = 7 度量
    matrix: dict = {}
    for r in rows:
        cid = r[0]
        ym = r[1]
        if cid not in matrix:
            matrix[cid] = {}
        matrix[cid][ym] = {
            "begin_balance": float(r[2] or 0),
            "avg_balance": float(r[3] or 0),
            "current_amount": float(r[4] or 0),
            "interest_rate": float(r[5] or 0),
            "interest_amount": float(r[6] or 0),
            "capital_ratio": float(r[7] or 0),
            "risk_weight": float(r[8] or 0),
        }

    # 4) 按 L1 大类聚合（path 首段 = /L1_xxx/）
    categories: dict = {}
    node_by_id = {n["coa_node_id"]: n for n in nodes}
    for cid, ym_map in matrix.items():
        n = node_by_id.get(cid)
        if not n or not n["path"]:
            continue
        cat = n["path"].split("/")[1] if "/" in n["path"] else "其他"
        cat = cat.replace("L1_", "")
        bucket = categories.setdefault(cat, {})
        for ym, m in ym_map.items():
            cb = bucket.setdefault(ym, {
                "begin_balance": 0.0, "avg_balance": 0.0, "current_amount": 0.0,
                "interest_rate": 0.0, "interest_amount": 0.0,
                "capital_ratio": 0.0, "risk_weight": 0.0,
                "account_count": 0,
            })
            cb["begin_balance"] += m["begin_balance"]
            cb["avg_balance"] += m["avg_balance"]
            cb["current_amount"] += m["current_amount"]
            cb["interest_amount"] += m["interest_amount"]
            cb["account_count"] += 1

    # 计算加权值
    for cat, ym_map in categories.items():
        for ym, m in ym_map.items():
            if m["avg_balance"] > 0:
                m["interest_rate"] = round(
                    sum(matrix.get(cid, {}).get(ym, {}).get("interest_rate", 0) *
                        matrix.get(cid, {}).get(ym, {}).get("avg_balance", 0)
                        for cid in matrix) / m["avg_balance"], 4)
                m["capital_ratio"] = round(
                    sum(matrix.get(cid, {}).get(ym, {}).get("capital_ratio", 0) *
                        matrix.get(cid, {}).get(ym, {}).get("avg_balance", 0)
                        for cid in matrix) / m["avg_balance"], 4)
                m["risk_weight"] = round(
                    sum(matrix.get(cid, {}).get(ym, {}).get("risk_weight", 0) *
                        matrix.get(cid, {}).get(ym, {}).get("avg_balance", 0)
                        for cid in matrix) / m["avg_balance"], 4)
            for k in ("begin_balance", "avg_balance", "current_amount", "interest_amount"):
                m[k] = round(m[k], 4)

    return {
        "scheme_id": scheme_id,
        "start_date": start_date,
        "end_date": end_date,
        "dates": dates,
        "nodes": nodes,
        "matrix": matrix,
        "categories": categories,
    }


@router.get("/by-scheme")
async def list_by_scheme(
    scheme_id: int = Query(...),
    data_date: str = Query(..., description="YYYY-MM-DD"),
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    """按账户册方案 + 数据日期聚合：每行 = 一个账户册节点 + 当月 7 个度量"""
    rows = db.execute(
        text(f"""SELECT b.id, b.coa_node_id, n.node_code, n.node_name,
                       n.node_level, n.node_type, n.path,
                       b.current_amount, b.begin_balance, b.avg_balance,
                       b.interest_rate, b.interest_amount,
                       b.capital_ratio, b.risk_weight,
                       {','.join('b.' + c for c in GAP_COLS)},
                       b.calc_note
                FROM prcp_data_balance b
                JOIN prcp_coa_node n ON n.id=b.coa_node_id
                WHERE b.data_date=:d AND b.is_deleted=0
                  AND n.scheme_id=:s
                ORDER BY n.path"""),
        {"d": data_date, "s": scheme_id},
    ).fetchall()

    items = []
    for r in rows:
        gaps = [float(r[14 + i]) for i in range(24)]
        items.append({
            "id": r[0], "coa_node_id": r[1],
            "node_code": r[2], "node_name": r[3],
            "node_level": r[4], "node_type": r[5], "path": r[6],
            "current_amount": float(r[7] or 0),
            "begin_balance": float(r[8] or 0),
            "avg_balance": float(r[9] or 0),
            "interest_rate": float(r[10] or 0),
            "interest_amount": float(r[11] or 0),
            "capital_ratio": float(r[12] or 0),
            "risk_weight": float(r[13] or 0),
            "gaps": gaps,
            "sum_24m": sum(gaps),
            "calc_note": r[38],
        })
    return {"scheme_id": scheme_id, "data_date": data_date,
            "items": items, "total": len(items)}


@router.get("/dates")
async def list_dates(
    scheme_id: Optional[int] = None,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    """返回 prcp_data_balance 中所有有数据的数据日期（按倒序），可选 scheme 过滤"""
    where = ["b.is_deleted=0"]
    params = {}
    if scheme_id:
        where.append("n.scheme_id=:s")
        params["s"] = scheme_id
    rows = db.execute(
        text(f"""SELECT b.data_date, COUNT(*) AS cnt,
                       SUM(b.current_amount) AS total_amt,
                       SUM(b.interest_amount) AS total_int
                FROM prcp_data_balance b
                JOIN prcp_coa_node n ON n.id=b.coa_node_id
                WHERE {' AND '.join(where)}
                GROUP BY b.data_date
                ORDER BY b.data_date DESC"""),
        params,
    ).fetchall()
    items = [{
        "data_date": r[0].isoformat() if r[0] else None,
        "record_count": int(r[1]),
        "total_amount": float(r[2] or 0),
        "total_interest": float(r[3] or 0),
    } for r in rows]
    return {"items": items, "total": len(items)}


@router.get("/category-summary")
async def category_summary(
    data_date: str = Query(..., description="YYYY-MM-DD"),
    scheme_id: Optional[int] = None,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    """按大类（资产/负债/表外）汇总当月数据"""
    where = ["b.is_deleted=0", "b.data_date=:d", "n.node_level=3"]
    params = {"d": data_date}
    if scheme_id:
        where.append("n.scheme_id=:s")
        params["s"] = scheme_id

    # 先把所有 L3 节点及其 path 抓出来，然后按 L1 大类名（path 首段）聚合
    rows = db.execute(
        text(f"""SELECT n.id, n.node_code, n.node_name, n.path,
                       b.current_amount, b.begin_balance, b.avg_balance,
                       b.interest_rate, b.interest_amount,
                       b.capital_ratio, b.risk_weight,
                       {','.join('b.' + c for c in GAP_COLS)}
                FROM prcp_data_balance b
                JOIN prcp_coa_node n ON n.id=b.coa_node_id
                WHERE {' AND '.join(where)}"""),
        params,
    ).fetchall()

    # 按 L1 分类（path 第 1 段：/L1_xxx/）
    buckets: dict = {}
    for r in rows:
        path = r[3] or ""
        cat = path.split('/')[1] if '/' in path else '其他'
        cat = cat.replace('L1_', '')
        b = buckets.setdefault(cat, {
            "category": cat, "account_count": 0,
            "total_amount": 0.0, "total_begin": 0.0, "total_avg": 0.0,
            "total_interest": 0.0, "weighted_rate_sum": 0.0, "weighted_cap_sum": 0.0,
            "weighted_rw_sum": 0.0, "sum_24m": 0.0,
        })
        b["account_count"] += 1
        amt = float(r[4] or 0)
        avg = float(r[6] or 0)
        b["total_amount"] += amt
        b["total_begin"] += float(r[5] or 0)
        b["total_avg"] += avg
        b["total_interest"] += float(r[8] or 0)
        b["weighted_rate_sum"] += float(r[7] or 0) * avg
        b["weighted_cap_sum"] += float(r[9] or 0) * avg
        b["weighted_rw_sum"] += float(r[10] or 0) * avg
        gaps = [float(r[11 + i]) for i in range(24)]
        b["sum_24m"] += sum(gaps)

    items = []
    for b in buckets.values():
        avg = b["total_avg"] or 1
        items.append({
            "category": b["category"],
            "account_count": b["account_count"],
            "total_amount": round(b["total_amount"], 2),
            "total_begin": round(b["total_begin"], 2),
            "total_avg": round(b["total_avg"], 2),
            "total_interest": round(b["total_interest"], 2),
            "weighted_rate": round(b["weighted_rate_sum"] / avg, 4),
            "weighted_capital": round(b["weighted_cap_sum"] / avg, 4),
            "weighted_risk_weight": round(b["weighted_rw_sum"] / avg, 4),
            "sum_24m": round(b["sum_24m"], 2),
        })
    return {"data_date": data_date, "items": items, "total": len(items)}


@router.get("/export-xlsx")
async def export_xlsx(
    scheme_id: int = Query(...),
    start_date: str = Query(..., description="YYYY-MM-DD 起始月"),
    end_date: str = Query(..., description="YYYY-MM-DD 结束月"),
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    """导出账户册矩阵为 Excel（月份 × 7 指标列）"""
    scheme_row = db.execute(
        text("SELECT scheme_code, scheme_name FROM prcp_coa_scheme WHERE id=:s"),
        {"s": scheme_id},
    ).first()
    if not scheme_row:
        raise HTTPException(404, "账户册方案不存在")
    scheme_code, scheme_name = scheme_row[0], scheme_row[1]

    node_rows = db.execute(
        text("""SELECT id, node_code, node_name, parent_id, node_level, sort_order, description
                FROM prcp_coa_node
                WHERE scheme_id=:s AND is_deleted=0
                ORDER BY sort_order, path"""),
        {"s": scheme_id},
    ).fetchall()

    # 月份列表
    start = datetime.strptime(start_date, "%Y-%m-%d").replace(day=1)
    end = datetime.strptime(end_date, "%Y-%m-%d").replace(day=1)
    dates: List[str] = []
    cur = start
    while cur <= end:
        dates.append(cur.strftime("%Y-%m"))
        if cur.month == 12:
            cur = cur.replace(year=cur.year + 1, month=1)
        else:
            cur = cur.replace(month=cur.month + 1)

    # 取数据
    balance_rows = db.execute(
        text("""SELECT coa_node_id, data_date, current_amount, begin_balance, avg_balance,
                       interest_rate, interest_amount, capital_ratio, risk_weight
                FROM prcp_data_balance
                WHERE data_date BETWEEN :sd AND :ed AND is_deleted=0"""),
        {"sd": start_date, "ed": end_date},
    ).fetchall()

    matrix: dict = {}
    for r in balance_rows:
        nid, dt = r[0], r[1]
        ym = dt.strftime("%Y-%m") if hasattr(dt, "strftime") else str(dt)[:7]
        matrix.setdefault(nid, {})[ym] = {
            "current_amount": float(r[2] or 0),
            "begin_balance": float(r[3] or 0),
            "avg_balance": float(r[4] or 0),
            "interest_rate": float(r[5] or 0),
            "interest_amount": float(r[6] or 0),
            "capital_ratio": float(r[7] or 0),
            "risk_weight": float(r[8] or 0),
        }

    # 生成 xlsx
    wb = Workbook()
    ws = wb.active
    ws.title = f"账户册矩阵"

    bold = Font(bold=True)
    blue_bold = Font(bold=True, color="1D39C4")
    header_fill = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
    center = Alignment(horizontal="center", vertical="center")

    # 标题
    title_cell = ws.cell(row=1, column=1,
        value=f"账户册总表（{scheme_name}）—— 余额：亿元；利率、资本占用比例、风险权重：%；平均利息收支：亿元/月")
    title_cell.font = bold
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4 + len(dates) * 7)

    # 一级表头：月份（合并 7 列）
    for col_idx, h in enumerate(["账户册编码", "账户册名称", "大类", "业务口径说明"], 1):
        c = ws.cell(row=2, column=col_idx, value=h)
        c.font = bold
        c.fill = header_fill
        c.alignment = center
    for ym_idx, ym in enumerate(dates):
        start_col = 5 + ym_idx * 7
        ws.merge_cells(start_row=2, start_column=start_col, end_row=2, end_column=start_col + 6)
        mc = ws.cell(row=2, column=start_col, value=ym)
        mc.font = blue_bold
        mc.alignment = center

    # 二级表头：7 指标
    for ym_idx in range(len(dates)):
        start_col = 5 + ym_idx * 7
        for m_idx, m in enumerate(EXPORT_MEASURES):
            c = ws.cell(row=3, column=start_col + m_idx, value=m["name"])
            c.font = bold
            c.alignment = center

    # 数据行
    # 父子映射
    parent_map = {n[0]: n[3] for n in node_rows}  # id -> parent_id
    id_to_name = {n[0]: n[2] for n in node_rows}

    def resolve_cat(node_id: int) -> str:
        """递归找到 L1 大类名"""
        cur = node_id
        while cur and parent_map.get(cur):
            p = parent_map[cur]
            # 检查 p 自己是不是 L1
            for n in node_rows:
                if n[0] == p and n[4] == 1:
                    return n[2]
            cur = p
        return ""

    for row_idx, n in enumerate(node_rows, 4):
        nid, code, name, parent_id, level, sort_o, desc = n
        cat = ""
        if level == 1:
            cat = name
        elif level in (2, 3):
            cat = resolve_cat(nid)
        ws.cell(row=row_idx, column=1, value=code)
        ws.cell(row=row_idx, column=2, value=name)
        ws.cell(row=row_idx, column=3, value=cat)
        ws.cell(row=row_idx, column=4, value=desc or "")
        for ym_idx, ym in enumerate(dates):
            start_col = 5 + ym_idx * 7
            cell = matrix.get(nid, {}).get(ym, {})
            for m_idx, m in enumerate(EXPORT_MEASURES):
                v = cell.get(m["key"])
                if v is not None:
                    ws.cell(row=row_idx, column=start_col + m_idx, value=v)

    # 列宽
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 28
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 30
    for ym_idx in range(len(dates)):
        for m_idx in range(7):
            col_letter = get_column_letter(5 + ym_idx * 7 + m_idx)
            ws.column_dimensions[col_letter].width = 12

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"prcp_balance_{scheme_code}_{start.strftime('%Y%m')}-{end.strftime('%Y%m')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


@router.post("/import-xlsx")
async def import_xlsx(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    """从 Excel 批量导入账户册月度数据
    Excel 格式：第 1 行标题；第 2 行月份（一级，合并 7 列）；第 3 行指标（二级）；第 4 行起数据
    列：A=账户册编码 B=账户册名称 C=大类 D=业务口径说明 E起=每个月 7 列
    """
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(400, "仅支持 .xlsx 格式文件")

    content = await file.read()
    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active

    # 读第 2 行（合并格：月份）
    header_row2 = [(c.value, c.column) for c in ws[2] if c.value]

    # 解析月份列范围（合并格的左上角单元格 = 月份值）
    ym_starts: list = []
    for v, col in header_row2:
        if v and re.match(r"^\d{4}-\d{2}$", str(v)):
            ym_starts.append((str(v), col))

    if not ym_starts:
        raise HTTPException(400, "Excel 第 2 行缺少月份表头（格式 YYYY-MM）")

    inserted = 0
    updated = 0
    skipped = 0
    errors: list = []
    uid = user.get("id", 1) if isinstance(user, dict) else getattr(user, "id", 1)

    # 预加载所有 node_code → id
    code_to_id = {}
    for r in db.execute(
        text("SELECT id, node_code FROM prcp_coa_node WHERE is_deleted=0")
    ).fetchall():
        code_to_id[r[1]] = r[0]

    for row in ws.iter_rows(min_row=4, values_only=False):
        code_cell = row[0].value
        if not code_cell:
            continue
        code = str(code_cell).strip()
        # 跳过 L1/L2 节点（编码以 L 开头）
        if code.startswith("L") and "_" in code:
            skipped += 1
            continue
        node_id = code_to_id.get(code)
        if not node_id:
            errors.append(f"编码 {code} 不存在")
            continue

        for ym, start_col in ym_starts:
            data = {}
            for m_idx, m in enumerate(EXPORT_MEASURES):
                col_idx = start_col + m_idx - 1  # 0-indexed
                if col_idx < len(row):
                    val = row[col_idx].value
                    if val is not None and val != "":
                        try:
                            data[m["key"]] = float(val)
                        except (TypeError, ValueError):
                            pass
            if not data:
                continue

            data_date = f"{ym}-01"
            existing = db.execute(
                text("SELECT id FROM prcp_data_balance WHERE coa_node_id=:n AND data_date=:d AND is_deleted=0"),
                {"n": node_id, "d": data_date},
            ).first()
            if existing:
                set_clauses = ", ".join(f"{k}=:{k}" for k in data)
                params = {**data, "id": existing[0], "u": uid}
                db.execute(
                    text(f"UPDATE prcp_data_balance SET {set_clauses}, updated_by=:u WHERE id=:id"),
                    params,
                )
                updated += 1
            else:
                cols = ["coa_node_id", "data_date"] + list(data.keys()) + ["created_by", "updated_by"]
                placeholders = ", ".join(f":{c}" for c in cols)
                params = {**data, "n": node_id, "d": data_date, "u": uid}
                db.execute(
                    text(f"INSERT INTO prcp_data_balance ({', '.join(cols)}) VALUES ({placeholders})"),
                    params,
                )
                inserted += 1

    db.commit()
    return {
        "inserted": inserted,
        "updated": updated,
        "skipped": skipped,
        "errors": errors[:20],
        "total_errors": len(errors),
    }
