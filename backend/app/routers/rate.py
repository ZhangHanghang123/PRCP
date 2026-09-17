"""利率管理 API — 收益率曲线方案 + 利率点矩阵
每条曲线 × 每个数据日期 = 一行 prcp_rate_point
13 个期限利率作为列存储（rate_d1..rate_y30）
"""
import io
import urllib.parse
from typing import Optional, List
from datetime import datetime, date
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import text
from sqlalchemy.orm import Session
from pydantic import BaseModel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from app.database import get_db
from app.auth import get_current_user

router = APIRouter(prefix="/rate", tags=["利率管理"])

# 13 个期限点（统一两端：短端 + 长端）
TERM_KEYS = [
    "d1", "d7", "m1", "m3", "m6",
    "y1", "y2", "y3", "y5",
    "y10", "y15", "y20", "y30",
]
TERM_NAMES = {
    "d1": "1日", "d7": "7日",
    "m1": "1M", "m3": "3M", "m6": "6M",
    "y1": "1Y", "y2": "2Y", "y3": "3Y", "y5": "5Y",
    "y10": "10Y", "y15": "15Y", "y20": "20Y", "y30": "30Y",
}
# 关键期限（前端高亮）
KEY_TERMS = ["y1", "y5", "y10"]


# ============================================================
# 1. 曲线方案管理
# ============================================================
class SchemeIn(BaseModel):
    curve_code: str
    curve_name: str
    curve_type: str
    ccy: str = "CNY"
    data_source: str = "WIND"
    description: Optional[str] = None
    status: str = "ACTIVE"


@router.get("/schemes")
async def list_schemes(
    curve_type: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    """列出所有曲线方案（含最新利率点统计）"""
    where = ["s.is_deleted=0"]
    params = {}
    if curve_type:
        where.append("s.curve_type = :t")
        params["t"] = curve_type
    if status:
        where.append("s.status = :st")
        params["st"] = status
    rows = db.execute(text(f"""
        SELECT s.id, s.curve_code, s.curve_name, s.curve_type, s.ccy, s.data_source,
               s.description, s.status, s.created_at, s.updated_at,
               (SELECT COUNT(*) FROM prcp_rate_point p
                WHERE p.curve_id=s.id AND p.is_deleted=0) AS point_count,
               (SELECT MAX(data_date) FROM prcp_rate_point
                WHERE curve_id=s.id AND is_deleted=0) AS latest_date,
               (SELECT rate_y10 FROM prcp_rate_point
                WHERE curve_id=s.id AND is_deleted=0
                ORDER BY data_date DESC LIMIT 1) AS latest_y10
        FROM prcp_rate_scheme s
        WHERE {' AND '.join(where)}
        ORDER BY s.curve_type, s.curve_code
    """), params).fetchall()
    items = []
    for r in rows:
        items.append({
            "id": r[0], "curve_code": r[1], "curve_name": r[2],
            "curve_type": r[3], "ccy": r[4], "data_source": r[5],
            "description": r[6], "status": r[7],
            "created_at": r[8].isoformat() if r[8] else None,
            "updated_at": r[9].isoformat() if r[9] else None,
            "point_count": int(r[10] or 0),
            "latest_date": str(r[11]) if r[11] else None,
            "latest_y10": float(r[12]) if r[12] is not None else None,
        })
    return {"items": items, "total": len(items)}


@router.post("/schemes")
async def create_scheme(p: SchemeIn, db: Session = Depends(get_db),
                         user=Depends(get_current_user)):
    """新建曲线方案"""
    uid = user.get("id") if isinstance(user, dict) else getattr(user, "id", 1)
    try:
        r = db.execute(text("""
            INSERT INTO prcp_rate_scheme
              (curve_code, curve_name, curve_type, ccy, data_source,
               description, status, created_by, updated_by)
            VALUES (:c, :n, :t, :cy, :ds, :d, :s, :u, :u)
        """), {"c": p.curve_code, "n": p.curve_name, "t": p.curve_type,
               "cy": p.ccy, "ds": p.data_source, "d": p.description,
               "s": p.status, "u": uid})
        db.commit()
        return {"id": r.lastrowid, "curve_code": p.curve_code, "ok": True}
    except Exception as e:
        db.rollback()
        raise HTTPException(400, f"创建失败：{e}")


@router.put("/schemes/{sid}")
async def update_scheme(sid: int, p: SchemeIn, db: Session = Depends(get_db),
                         user=Depends(get_current_user)):
    """更新曲线方案"""
    uid = user.get("id") if isinstance(user, dict) else getattr(user, "id", 1)
    r = db.execute(text("""
        UPDATE prcp_rate_scheme SET
          curve_name=:n, curve_type=:t, ccy=:cy, data_source=:ds,
          description=:d, status=:s, updated_by=:u
        WHERE id=:id AND is_deleted=0
    """), {"n": p.curve_name, "t": p.curve_type, "cy": p.ccy,
           "ds": p.data_source, "d": p.description, "s": p.status,
           "u": uid, "id": sid})
    if r.rowcount == 0:
        raise HTTPException(404, "曲线方案不存在")
    db.commit()
    return {"ok": True}


@router.delete("/schemes/{sid}")
async def delete_scheme(sid: int, db: Session = Depends(get_db),
                         user=Depends(get_current_user)):
    """软删除曲线方案（同时软删除其所有利率点）"""
    r = db.execute(text("""
        UPDATE prcp_rate_scheme SET is_deleted=1, updated_by=:u
        WHERE id=:id AND is_deleted=0
    """), {"u": user.get("id") if isinstance(user, dict) else 1, "id": sid})
    if r.rowcount == 0:
        raise HTTPException(404, "曲线方案不存在")
    db.execute(text("""
        UPDATE prcp_rate_point SET is_deleted=1
        WHERE curve_id=:id AND is_deleted=0
    """), {"id": sid})
    db.commit()
    return {"ok": True}


# ============================================================
# 2. 利率点管理（核心）
# ============================================================
class PointIn(BaseModel):
    curve_code: str
    data_date: str  # YYYY-MM-DD
    ccy: str = "CNY"
    rates: dict = {}  # {"d1": 1.5, "d7": 1.65, ..., "y30": 3.05}
    source_date: Optional[str] = None
    remark: Optional[str] = None


@router.get("/points")
async def list_points(
    curve_code: Optional[str] = Query(None),
    data_date: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    """列出利率点"""
    where = ["p.is_deleted=0"]
    params = {}
    if curve_code:
        where.append("p.curve_code = :c")
        params["c"] = curve_code
    if data_date:
        where.append("p.data_date = :d")
        params["d"] = data_date
    rows = db.execute(text(f"""
        SELECT p.id, p.curve_id, p.curve_code, s.curve_name, p.data_date, p.ccy,
               p.rate_d1, p.rate_d7, p.rate_m1, p.rate_m3, p.rate_m6,
               p.rate_y1, p.rate_y2, p.rate_y3, p.rate_y5, p.rate_y10,
               p.rate_y15, p.rate_y20, p.rate_y30,
               p.curve_shift_bps, p.curve_slope, p.source_date, p.remark,
               p.created_at, p.updated_at
        FROM prcp_rate_point p
        JOIN prcp_rate_scheme s ON s.id=p.curve_id
        WHERE {' AND '.join(where)}
        ORDER BY p.curve_code, p.data_date DESC
    """), params).fetchall()
    items = []
    for r in rows:
        items.append({
            "id": r[0], "curve_id": r[1], "curve_code": r[2], "curve_name": r[3],
            "data_date": str(r[4]), "ccy": r[5],
            "rates": {
                "d1": float(r[6] or 0), "d7": float(r[7] or 0),
                "m1": float(r[8] or 0), "m3": float(r[9] or 0), "m6": float(r[10] or 0),
                "y1": float(r[11] or 0), "y2": float(r[12] or 0),
                "y3": float(r[13] or 0), "y5": float(r[14] or 0),
                "y10": float(r[15] or 0), "y15": float(r[16] or 0),
                "y20": float(r[17] or 0), "y30": float(r[18] or 0),
            },
            "curve_shift_bps": float(r[19] or 0),
            "curve_slope": float(r[20] or 0),
            "source_date": str(r[21]) if r[21] else None,
            "remark": r[22],
            "created_at": r[23].isoformat() if r[23] else None,
            "updated_at": r[24].isoformat() if r[24] else None,
        })
    return {"items": items, "total": len(items)}


@router.post("/points")
async def upsert_point(p: PointIn, db: Session = Depends(get_db),
                       user=Depends(get_current_user)):
    """新增/更新一条利率点（按 curve_code + data_date 唯一）"""
    uid = user.get("id") if isinstance(user, dict) else getattr(user, "id", 1)
    # 1. 查 curve_id
    scheme_row = db.execute(text("""
        SELECT id FROM prcp_rate_scheme WHERE curve_code=:c AND is_deleted=0
    """), {"c": p.curve_code}).first()
    if not scheme_row:
        raise HTTPException(404, f"曲线方案 {p.curve_code} 不存在")
    curve_id = scheme_row[0]

    # 2. 准备 13 个利率值
    rate_values = {k: float(p.rates.get(k, 0) or 0) for k in TERM_KEYS}
    # 3. 自动计算 curve_slope = 10Y - 1Y
    slope = rate_values["y10"] - rate_values["y1"]

    # 4. 计算 BP 平移（vs 上一个数据日期的 10Y）
    prev = db.execute(text("""
        SELECT rate_y10 FROM prcp_rate_point
        WHERE curve_code=:c AND data_date < :d AND is_deleted=0
        ORDER BY data_date DESC LIMIT 1
    """), {"c": p.curve_code, "d": p.data_date}).first()
    if prev:
        shift_bps = (rate_values["y10"] - float(prev[0])) * 100
    else:
        shift_bps = 0.0

    # 5. UPSERT（INSERT ON DUPLICATE KEY UPDATE）
    try:
        db.execute(text("""
            INSERT INTO prcp_rate_point
              (curve_id, curve_code, data_date, ccy,
               rate_d1, rate_d7, rate_m1, rate_m3, rate_m6,
               rate_y1, rate_y2, rate_y3, rate_y5, rate_y10,
               rate_y15, rate_y20, rate_y30,
               curve_shift_bps, curve_slope, source_date, remark,
               created_by, updated_by)
            VALUES
              (:cid, :c, :d, :cy,
               :d1, :d7, :m1, :m3, :m6,
               :y1, :y2, :y3, :y5, :y10,
               :y15, :y20, :y30,
               :shift, :slope, :sd, :r,
               :u, :u)
            ON DUPLICATE KEY UPDATE
              ccy=:cy,
              rate_d1=:d1, rate_d7=:d7, rate_m1=:m1, rate_m3=:m3, rate_m6=:m6,
              rate_y1=:y1, rate_y2=:y2, rate_y3=:y3, rate_y5=:y5, rate_y10=:y10,
              rate_y15=:y15, rate_y20=:y20, rate_y30=:y30,
              curve_shift_bps=:shift, curve_slope=:slope,
              source_date=:sd, remark=:r, updated_by=:u,
              is_deleted=0
        """), {
            "cid": curve_id, "c": p.curve_code, "d": p.data_date, "cy": p.ccy,
            "d1": rate_values["d1"], "d7": rate_values["d7"],
            "m1": rate_values["m1"], "m3": rate_values["m3"], "m6": rate_values["m6"],
            "y1": rate_values["y1"], "y2": rate_values["y2"],
            "y3": rate_values["y3"], "y5": rate_values["y5"],
            "y10": rate_values["y10"],
            "y15": rate_values["y15"], "y20": rate_values["y20"], "y30": rate_values["y30"],
            "shift": shift_bps, "slope": slope,
            "sd": p.source_date or p.data_date, "r": p.remark,
            "u": uid,
        })
        db.commit()
        return {
            "ok": True,
            "curve_code": p.curve_code,
            "data_date": p.data_date,
            "curve_shift_bps": round(shift_bps, 2),
            "curve_slope": round(slope, 4),
        }
    except Exception as e:
        db.rollback()
        raise HTTPException(400, f"保存失败：{e}")


@router.delete("/points/{pid}")
async def delete_point(pid: int, db: Session = Depends(get_db),
                       user=Depends(get_current_user)):
    """软删除利率点"""
    r = db.execute(text("""
        UPDATE prcp_rate_point SET is_deleted=1
        WHERE id=:id AND is_deleted=0
    """), {"id": pid})
    if r.rowcount == 0:
        raise HTTPException(404, "利率点不存在")
    db.commit()
    return {"ok": True}


# ============================================================
# 3. 历史曲线对比
# ============================================================
@router.get("/compare")
async def compare_curves(
    curve_code: str = Query(...),
    start_date: Optional[str] = Query(None),
    end_date: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    """历史曲线对比（按日期序列）"""
    where = ["curve_code=:c", "is_deleted=0"]
    params = {"c": curve_code}
    if start_date:
        where.append("data_date >= :s")
        params["s"] = start_date
    if end_date:
        where.append("data_date <= :e")
        params["e"] = end_date
    rows = db.execute(text(f"""
        SELECT data_date,
               rate_d1, rate_d7, rate_m1, rate_m3, rate_m6,
               rate_y1, rate_y2, rate_y3, rate_y5, rate_y10,
               rate_y15, rate_y20, rate_y30, curve_slope
        FROM prcp_rate_point
        WHERE {' AND '.join(where)}
        ORDER BY data_date ASC
    """), params).fetchall()
    items = []
    for r in rows:
        items.append({
            "data_date": str(r[0]),
            "rates": {
                "d1": float(r[1] or 0), "d7": float(r[2] or 0),
                "m1": float(r[3] or 0), "m3": float(r[4] or 0), "m6": float(r[5] or 0),
                "y1": float(r[6] or 0), "y2": float(r[7] or 0),
                "y3": float(r[8] or 0), "y5": float(r[9] or 0),
                "y10": float(r[10] or 0),
                "y15": float(r[11] or 0), "y20": float(r[12] or 0), "y30": float(r[13] or 0),
            },
            "curve_slope": float(r[14] or 0),
        })
    return {"curve_code": curve_code, "items": items, "total": len(items)}


# ============================================================
# 4. 导出 Excel
# ============================================================
@router.get("/export-xlsx")
async def export_xlsx(
    curve_code: str = Query(...),
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    """导出曲线为 Excel（多日期 × 13 期限点矩阵）"""
    rows = db.execute(text("""
        SELECT data_date,
               rate_d1, rate_d7, rate_m1, rate_m3, rate_m6,
               rate_y1, rate_y2, rate_y3, rate_y5, rate_y10,
               rate_y15, rate_y20, rate_y30, curve_slope, remark
        FROM prcp_rate_point
        WHERE curve_code=:c AND is_deleted=0
        ORDER BY data_date DESC
    """), {"c": curve_code}).fetchall()
    if not rows:
        raise HTTPException(404, "曲线无利率点")

    # 拿曲线名称
    s = db.execute(text("""
        SELECT curve_name FROM prcp_rate_scheme WHERE curve_code=:c AND is_deleted=0
    """), {"c": curve_code}).first()
    curve_name = s[0] if s else curve_code

    wb = Workbook()
    ws = wb.active
    ws.title = curve_code

    # row 1: 标题
    ws.cell(1, 2, f"{curve_code} · {curve_name} · 收益率曲线明细")

    # row 2: 二级表头
    ws.cell(2, 3, "期限点（年化利率 %）")

    # row 3: 详细列名
    headers = ["数据日期"] + [TERM_NAMES[k] for k in TERM_KEYS] + ["斜率", "备注"]
    for col, h in enumerate(headers, start=1):
        ws.cell(3, col, h)

    # 样式
    FILL_KEY = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    FILL_HEAD = PatternFill(start_color="EDE9FE", end_color="EDE9FE", fill_type="solid")
    FONT_KEY = Font(bold=True, color="92400E")

    # row 4+: 数据
    for ri, r in enumerate(rows, start=4):
        ws.cell(ri, 1, str(r[0]))
        for i, k in enumerate(TERM_KEYS, start=2):
            v = float(r[i - 1] or 0)
            cell = ws.cell(ri, i, round(v, 4))
            if k in KEY_TERMS:
                cell.fill = FILL_KEY
                if k == "y10":
                    cell.font = FONT_KEY
        ws.cell(ri, 15, float(r[14] or 0))  # slope
        ws.cell(ri, 16, r[15] or "")  # remark

    # 表头样式
    for col in range(1, len(headers) + 1):
        ws.cell(3, col).fill = FILL_HEAD
        ws.cell(3, col).font = Font(bold=True)
        ws.cell(3, col).alignment = Alignment(horizontal="center")

    # 列宽
    ws.column_dimensions["A"].width = 14
    for i, k in enumerate(TERM_KEYS, start=2):
        ws.column_dimensions[chr(64 + i)].width = 10
    ws.column_dimensions["O"].width = 10
    ws.column_dimensions["P"].width = 30
    ws.freeze_panes = "B4"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"利率曲线_{curve_code}.xlsx"
    quoted = urllib.parse.quote(fname)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename=\"rate.xlsx\"; filename*=UTF-8''{quoted}",
        },
    )


# ============================================================
# 5. 工具：按期限点快速查询（供其他模块调用）
# ============================================================
@router.get("/lookup")
async def lookup_rate(
    curve_code: str = Query(...),
    data_date: str = Query(...),
    term: str = Query(..., description="期限点：d1/d7/m1/.../y30"),
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    """按曲线 + 日期 + 期限点查单一利率（供 KPI/模型/反算使用）"""
    if term not in TERM_KEYS:
        raise HTTPException(400, f"期限点 {term} 非法，应为 {TERM_KEYS}")
    row = db.execute(text(f"""
        SELECT rate_{term} FROM prcp_rate_point
        WHERE curve_code=:c AND data_date=:d AND is_deleted=0
        LIMIT 1
    """), {"c": curve_code, "d": data_date}).first()
    if not row:
        raise HTTPException(404, f"无 {curve_code} @ {data_date} 利率数据")
    return {"curve_code": curve_code, "data_date": data_date,
            "term": term, "rate": float(row[0] or 0)}