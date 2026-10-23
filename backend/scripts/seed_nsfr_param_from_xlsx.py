# -*- coding: utf-8 -*-
"""从 docs/v2/NSFR.xlsx 导入初始 NSFR 参数到 prcp_nsfr_param
- 方案：ZX_COA（scheme_id=1, 共 57 节点）
- 数据日期：2025/12/31
"""
import sys
import os
from urllib.parse import quote_plus
import openpyxl
from sqlalchemy import create_engine, text

DB_URL = f"mysql+pymysql://prcp:{quote_plus('Prcp@2026')}@127.0.0.1:3306/prcp_db?charset=utf8mb4"
NSFR_FILE = os.environ.get('NSFR_FILE', r'C:\银行经营\PRCP\docs\v2\NSFR.xlsx')
if not os.path.exists(NSFR_FILE):
    linux_path = '/home/almd/prcp/docs/v2/NSFR.xlsx'
    if os.path.exists(linux_path):
        NSFR_FILE = linux_path
SCHEME_CODE = "ZX_COA"
SCHEME_ID = 1
DATA_DATE = "2025-12-31"


def _to_int(v):
    """是/否 → 1/0"""
    if v is None:
        return 0
    s = str(v).strip()
    if s in ("是", "yes", "Y", "y", "1", "true", "True"):
        return 1
    return 0


def _to_num(v):
    """任意 → 数字（None → 0）"""
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if s in ("-", ""):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def main():
    wb = openpyxl.load_workbook(NSFR_FILE, data_only=True)
    if 'NSFR项目明细' not in wb.sheetnames:
        print(f"❌ {NSFR_FILE} 没有 'NSFR项目明细' sheet")
        sys.exit(1)
    ws = wb['NSFR项目明细']

    engine = create_engine(DB_URL)
    with engine.connect() as conn:
        # 1. 加载 ZX_COA 节点映射
        node_rows = conn.execute(text("""
            SELECT id, node_code, node_name FROM prcp_coa_node
            WHERE scheme_id=:s AND is_deleted=0
        """), {"s": SCHEME_ID}).fetchall()
        node_map = {r[1]: (r[0], r[2]) for r in node_rows}
        print(f"节点映射: {len(node_map)} 条")

        # 2. 软删旧数据
        conn.execute(text("""
            UPDATE prcp_nsfr_param SET is_deleted=1, updated_by=1
            WHERE scheme_id=:s AND data_date=STR_TO_DATE(:d, '%Y-%m-%d')
        """), {"s": SCHEME_ID, "d": DATA_DATE})
        print(f"已软删旧数据")

        # 3. 导入 NSFR.xlsx
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        ok = skip = 0
        for row in rows:
            code = row[0]
            if not code or not isinstance(code, str):
                continue
            if not code.startswith('S'):
                continue
            if code not in node_map:
                print(f"  [跳过] {code} 不在 ZX_COA 节点表")
                skip += 1
                continue
            nid, nname = node_map[code]
            is_asf_v = _to_int(row[4])
            asf_f = _to_num(row[5])
            asf_op = row[6] or '+'
            is_rsf_v = _to_int(row[7])
            rsf_f = _to_num(row[8])
            rsf_op = row[9] or '+'
            note = (row[12] or '')[:500]
            new_id = f"{SCHEME_CODE}_{code}_{DATA_DATE.replace('-', '')}"

            conn.execute(text("""
                INSERT INTO prcp_nsfr_param
                  (id, scheme_id, scheme_code, node_id, node_code, node_name, data_date,
                   is_asf, asf_factor, asf_operator,
                   is_rsf, rsf_factor, rsf_operator,
                   rule_note, status, created_by, updated_by)
                VALUES
                  (:id, :sid, :sc, :nid, :nc, :nn, STR_TO_DATE(:dd, '%Y-%m-%d'),
                   :is_a, :a_f, :a_op,
                   :is_r, :r_f, :r_op,
                   :note, 'ACTIVE', 1, 1)
            """), {
                "id": new_id, "sid": SCHEME_ID, "sc": SCHEME_CODE,
                "nid": nid, "nc": code, "nn": nname,
                "dd": DATA_DATE,
                "is_a": is_asf_v, "a_f": asf_f, "a_op": asf_op,
                "is_r": is_rsf_v, "r_f": rsf_f, "r_op": rsf_op,
                "note": note,
            })
            ok += 1

        conn.commit()
        print(f"\n[完成] 导入 {ok} 条，跳过 {skip} 条")

    # 4. 验证
    with engine.connect() as conn:
        stat = conn.execute(text("""
            SELECT
              COUNT(*) AS total,
              SUM(is_asf) AS asf_yes,
              SUM(is_rsf) AS rsf_yes
            FROM prcp_nsfr_param
            WHERE is_deleted=0 AND data_date=STR_TO_DATE(:d, '%Y-%m-%d')
        """), {"d": DATA_DATE}).fetchone()
        print(f"\n=== 验证 ===")
        print(f"总数: {stat[0]}, ASF 节点: {stat[1]}, RSF 节点: {stat[2]}")


if __name__ == "__main__":
    main()