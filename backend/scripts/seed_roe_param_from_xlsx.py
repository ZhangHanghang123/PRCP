"""从 ROE.xlsx 导入 ROE 参数补录到 prcp_roe_param
按 ZX_COA 方案过滤：xlsx 中 ZX_COA 不存在的节点自动跳过

用法（服务器端）：
    venv/bin/python3 scripts/seed_roe_from_xlsx.py
"""
import os
import sys
from urllib.parse import quote_plus

import openpyxl
from sqlalchemy import create_engine, text

DB_URL = f"mysql+pymysql://prcp:{quote_plus('Prcp@2026')}@127.0.0.1:3306/prcp_db?charset=utf8mb4"
SCHEME_ID = 1
SCHEME_CODE = "ZX_COA"
DATA_DATE = "2025-12-31"

XLSX_PATHS = [
    os.environ.get("ROE_FILE", ""),
    r"C:\银行经营\PRCP\docs\v2\ROE.xlsx",
    r"/home/almd/prcp/docs/v2/ROE.xlsx",
]
XLSX_FILE = next((p for p in XLSX_PATHS if p and os.path.exists(p)), None)
if not XLSX_FILE:
    print(f"  找不到 ROE.xlsx（请检查路径或环境变量 ROE_FILE）")
    sys.exit(1)
print(f"  ROE.xlsx 路径: {XLSX_FILE}")

engine = create_engine(DB_URL, echo=False)
def extract_fn(row, code, node_id, node_name):
    """ROE：净利润+净资产（符号+系数）"""
    is_a = str(row[4]).strip() if row[4] else ""
    a_symbol = str(row[5]).strip() if row[5] else "+"
    a_factor = float(row[6]) if isinstance(row[6], (int, float)) else 0.0
    is_b = str(row[7]).strip() if row[7] else ""
    b_symbol = str(row[8]).strip() if row[8] else "+"
    b_factor = float(row[9]) if isinstance(row[9], (int, float)) else 0.0
    rule_note = (str(row[10]) if len(row) > 10 and row[10] else "").strip()
    cur_bal = float(row[3]) if isinstance(row[3], (int, float)) else 0.0

    is_np = 1 if is_a == "是" else 0
    is_na = 1 if is_b == "是" else 0

    a_symbol = "+" if a_symbol in ("+", "加", "") else "-"
    b_symbol = "+" if b_symbol in ("+", "加", "") else "-"

    return {
        "scheme_id": 1, "scheme_code": "ZX_COA",
        "node_id": node_id, "node_code": code, "node_name": node_name,
        "data_date": DATA_DATE,
        "is_net_profit": is_np,
        "net_profit_symbol": a_symbol, "net_profit_factor": a_factor, "net_profit_category": "",
        "is_net_asset": is_na,
        "net_asset_symbol": b_symbol, "net_asset_factor": b_factor, "net_asset_category": "",
        "current_balance": cur_bal, "rule_note": rule_note,
        "status": "ACTIVE", "is_deleted": 0, "created_by": "seed", "updated_by": "seed",
    }, [
        "scheme_id", "scheme_code", "node_id", "node_code", "node_name", "data_date",
        "is_net_profit", "net_profit_symbol", "net_profit_factor", "net_profit_category",
        "is_net_asset", "net_asset_symbol", "net_asset_factor", "net_asset_category",
        "current_balance", "rule_note", "status", "is_deleted", "created_by", "updated_by",
    ]

def main():
    # 1. 加载 xlsx
    print(f"\n=== 1. 读取 XLSX ===")
    wb = openpyxl.load_workbook(XLSX_FILE, data_only=True)
    ws = wb["ROE项目明细"]
    print(f"  Sheet 'ROE项目明细': {ws.max_row - 1} 行数据")

    # 2. 加载 ZX_COA 节点清单
    print("\n=== 2. 加载 ZX_COA 节点清单 ===")
    with engine.begin() as conn:
        nodes = conn.execute(text("""
            SELECT id, node_code, node_name
            FROM prcp_coa_node
            WHERE scheme_id=:sid AND is_deleted=0
        """), {"sid": SCHEME_ID}).fetchall()
    node_map = {row[1]: (row[0], row[2]) for row in nodes}
    print(f"  ZX_COA 节点数: {len(node_map)}")

    # 3. 解析 xlsx
    print("\n=== 3. 解析数据 ===")
    parsed = []
    skipped = []
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_idx == 1:
            continue
        if not row or not row[0]:
            continue
        code = str(row[0]).strip()
        if not code:
            continue
        if code not in node_map:
            skipped.append(code)
            continue
        node_id, node_name = node_map[code]
        parsed.append((code, node_id, node_name, row))
    print(f"  解析: {len(parsed)} 条；跳过 ZX_COA 不存在 {len(skipped)} 条")
    if skipped:
        print(f"    跳过节点: {skipped[:8]}{'...' if len(skipped)>8 else ''}")

    # 4. 清表 + 重插
    print("\n=== 4. 清表 + 重插到 prcp_roe_param ===")
    with engine.begin() as conn:
        conn.execute(text(f"DELETE FROM prcp_roe_param WHERE data_date=STR_TO_DATE(:dd, '%Y-%m-%d')"),
                     {"dd": DATA_DATE})
        inserted = 0
        for code, node_id, node_name, row in parsed:
            params, columns = extract_fn(row, code, node_id, node_name)
            if not params:
                continue
            params["id"] = f"{SCHEME_CODE}_{code}_{DATA_DATE.replace('-', '')}"
            cols = ", ".join(["id"] + columns)
            placeholders = ", ".join([f":{c}" for c in (["id"] + columns)])
            conn.execute(text(f"INSERT INTO prcp_roe_param ({cols}) VALUES ({placeholders})"), params)
            inserted += 1

    # 5. 统计
    print("\n=== 5. 统计 ===")
    with engine.begin() as conn:
        stat = conn.execute(text(f"""
            SELECT COUNT(*) total, SUM(is_net_profit) yes_a, SUM(is_net_asset) yes_b
            FROM prcp_roe_param WHERE is_deleted=0
        """)).fetchone()
    print(f"  导入: {inserted} 条 | 数据库当前: total={stat[0]}, A 类 yes={stat[1] or 0}, B 类 yes={stat[2] or 0}")
    print("\n✓ 完成")


if __name__ == "__main__":
    main()
