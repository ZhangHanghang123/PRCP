"""Test basic API after bucket refactor."""
import requests
import os

r = requests.post("https://wxfzhh.online/prcp/api/auth/login",
                   data={"username": "admin", "password": "admin123"}, timeout=10)
token = r.json()["access_token"]
H = {"Authorization": f"Bearer {token}"}

# 1. by-scheme-matrix
r = requests.get(
    "https://wxfzhh.online/prcp/api/basic/by-scheme-matrix",
    params={"scheme_id": 6, "data_date": "2025-12-31", "date_offset": 0, "offset_unit": "D"},
    headers=H,
)
print(f"matrix: status={r.status_code}")
d = r.json()
print(f"  matrix keys: {len(d['matrix'])}, categories: {list(d['categories'].keys())}")
m = d["matrix"][list(d["matrix"].keys())[0]]
print(f"  first node cid: {list(d['matrix'].keys())[0]}")
print(f"  has orig_m30? {('orig_m30' in m)}, orig_m30={m.get('orig_m30', 'N/A')}")
print(f"  has orig_y2? {('orig_y2' in m)}")
print(f"  has orig_y5? {('orig_y5' in m)}")
print(f"  has orig_y3? {('orig_y3' in m)}")
print(f"  has orig_m13? {('orig_m13' in m)}")
print(f"  has orig_m60? {('orig_m60' in m)}")

# 2. export-xlsx
r = requests.get(
    "https://wxfzhh.online/prcp/api/basic/export-xlsx",
    params={"scheme_id": 6, "data_date": "2025-12-31", "date_offset": 0, "offset_unit": "D"},
    headers=H,
)
print(f"\nexport-xlsx: status={r.status_code}, size={len(r.content)} bytes")
out = "C:/银行经营/PRCP/tmp_export/basic_2025-12-31_v3.xlsx"
os.makedirs(os.path.dirname(out), exist_ok=True)
with open(out, "wb") as f:
    f.write(r.content)
print(f"saved to {out}")

# 3. 校验 Excel
from openpyxl import load_workbook
wb = load_workbook(out)
ws = wb.active
print(f"\nExcel: max_row={ws.max_row}, max_col={ws.max_column}")
# row 3: 详细列名
hdrs = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
print(f"  col 1: {hdrs[0]}")
print(f"  col 11: {hdrs[10]}")
print(f"  col 12: {hdrs[11]} (orig 1)")
print(f"  col 14: {hdrs[13]} (orig 3M)")
print(f"  col 16: {hdrs[15]} (orig 6M)")
print(f"  col 17: {hdrs[16]} (orig 13M)")
print(f"  col 64: {hdrs[63]} (orig 60M)")
print(f"  col 65: {hdrs[64]} (orig 1Y)")
print(f"  col 69: {hdrs[68]} (orig 30Y)")
print(f"  col 70: {hdrs[69]} (rem 1日)")