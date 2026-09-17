"""Mock 数据测试 export 层级样式"""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

nodes = [
    (1,  "ZX_A001", "总资产",         1),
    (2,  "ZX_A002", "1.生息资产",     2),
    (3,  "ZX_A003", "（一）人民币",   3),
    (4,  "ZX_A007", "1.自营贷款",     4),
    (5,  "ZX_A011", "1.1.对公一般贷款", 5),
    (6,  "ZX_A028", "2.境内存放同业", 4),
    (7,  "ZX_L003", "总负债",         1),
    (8,  "ZX_L008", "（一）人民币小计", 2),
    (9,  "ZX_L031", "境内人民币自营存", 3),
    (10, "ZX_L032", "1.1.对公存款",    4),
    (11, "ZX_L033", "1.1.1.对公活期存款", 5),
    (12, "ZX_E005", "所有者权益",     1),
]

rows = {n[0]: (n[0], n[1], n[2], n[3], None, "A" if "A" in n[1] else ("L" if "L" in n[1] else "E"),
              1, *([0.0]*13), *([0.0]*13), "", None, 1000.0, 900.0, 0.04, 36.0, 0.75, "")
        for n in nodes}

wb = Workbook()
ws = wb.active
ws.title = "基础数据"

ws.cell(1, 3, "账户册总表（资产 / 负债 / 表外）")
ws.cell(2, 11, "原始期限(金额)")
ws.cell(2, 24, "剩余期限（金额）")

detail_headers = ["ID", "数据日期", "账户册编码", "账户册名称", "层级", "父级编码",
                  "是否末级", "大类", "偏移量", "偏移单位",
                  "1日","7日","1M","3M","6M","1Y","2Y","3Y","5Y","10Y","15Y","20Y","30Y",
                  "1日","7日","1M","3M","6M","1Y","2Y","3Y","5Y","10Y","15Y","20Y","30Y",
                  "ASF/RSF","HQLA","当前余额","平均余额","加权平均利率","平均利息收支","风险权重"]
for col, h in enumerate(detail_headers, start=1):
    ws.cell(3, col, h)

FILL_L1 = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
FILL_L2 = PatternFill(start_color="EAF1F8", end_color="EAF1F8", fill_type="solid")
FILL_SEP = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
FONT_L1 = Font(bold=True, size=11, color="1F4E78")
FONT_L2 = Font(bold=True, size=10)
FONT_NORMAL = Font(size=10)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")

prev_l1_code = None
ri = 4
for n in nodes:
    node_id, node_code, node_name, node_level = n

    if node_level == 1 and prev_l1_code is not None and node_code != prev_l1_code:
        for col in range(1, 44):
            cell = ws.cell(ri, col, "")
            cell.fill = FILL_SEP
        ri += 1
    if node_level == 1:
        prev_l1_code = node_code

    r = rows.get(node_id)
    ws.cell(ri, 1, f"2025-12-31-{node_code}")
    ws.cell(ri, 2, "2025-12-31")
    ws.cell(ri, 3, node_code)

    indent = "　" * max(node_level - 1, 0)
    name_cell = ws.cell(ri, 4, f"{indent}{node_name}")
    if node_level == 1:
        name_cell.font = FONT_L1; name_cell.fill = FILL_L1
    elif node_level == 2:
        name_cell.font = FONT_L2; name_cell.fill = FILL_L2
    else:
        name_cell.font = FONT_NORMAL
    name_cell.alignment = ALIGN_LEFT

    lv_cell = ws.cell(ri, 5, node_level)
    if node_level == 1:
        lv_cell.font = FONT_L1; lv_cell.fill = FILL_L1
    elif node_level == 2:
        lv_cell.font = FONT_L2; lv_cell.fill = FILL_L2
    lv_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.cell(ri, 6, r[4] if r and r[4] else "")
    ws.cell(ri, 7, r[6] if r and r[6] is not None else 0)
    ws.cell(ri, 8, r[5] if r and r[5] else "")
    ws.cell(ri, 9, 0)
    ws.cell(ri, 10, "D")

    if r:
        for i in range(13):
            ws.cell(ri, 11 + i, float(r[7 + i] or 0))
        for i in range(13):
            ws.cell(ri, 24 + i, float(r[20 + i] or 0))
        ws.cell(ri, 37, r[33] if r[33] is not None else "")
        ws.cell(ri, 38, float(r[34]) if r[34] is not None else "")
        ws.cell(ri, 39, float(r[35] or 0))
        ws.cell(ri, 40, float(r[36] or 0))
        ws.cell(ri, 41, float(r[37] or 0))
        ws.cell(ri, 42, float(r[38] or 0))
        ws.cell(ri, 43, float(r[39] or 0))

    if node_level == 1:
        for col in range(1, 44):
            cur = ws.cell(ri, col)
            if not cur.fill or cur.fill.start_color.rgb in (None, "00000000"):
                cur.fill = FILL_L1
    elif node_level == 2:
        for col in range(1, 44):
            cur = ws.cell(ri, col)
            if not cur.fill or cur.fill.start_color.rgb in (None, "00000000"):
                cur.fill = FILL_L2
    ri += 1

ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 12
ws.column_dimensions["C"].width = 12
ws.column_dimensions["D"].width = 40
ws.column_dimensions["E"].width = 8
ws.freeze_panes = "E4"

buf = io.BytesIO()
wb.save(buf)
buf.seek(0)
out_path = "/tmp/test_export_hier.xlsx"
with open(out_path, "wb") as f:
    f.write(buf.read())
print(f"OK saved {out_path}, size={len(buf.getvalue())} bytes, total rows used={ri-1}")