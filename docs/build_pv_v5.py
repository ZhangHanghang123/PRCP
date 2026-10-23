"""K 列现值 v5: 在 v4 基础上加回 B113-B176 折现因子公式

v4: K 列是数值, B113-B176 是硬编码数值
v5: K 列是数值, B113-B176 改回公式 (可审计 + 收益率变化时自动重算)
"""

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

EXCEL_PATH = r"C:/银行经营/PRCP/docs/v2/prcp_basic_20251231_v4.xlsx"

# ===================== 1. 加载 =====================
wb = load_workbook(EXCEL_PATH)
ws = wb['基础数据']

border = Border(
    left=Side(style="thin", color="BDBDBD"),
    right=Side(style="thin", color="BDBDBD"),
    top=Side(style="thin", color="BDBDBD"),
    bottom=Side(style="thin", color="BDBDBD"),
)

# ===================== 2. B113-B176 改回公式 =====================
# 月度: B113-B172 (m=1..60)
for m in range(1, 61):
    row = 112 + m
    rate_row = 47 + m  # B48..B107
    formula = f"=1/(1+$B${rate_row})^({m}/12)"
    ws.cell(row, 2, formula)
    ws.cell(row, 2).font = Font(name="微软雅黑", size=9, color="555555")
    ws.cell(row, 2).number_format = "0.000000000"
    ws.cell(row, 2).border = border
    ws.cell(row, 1, f"DF {m}M")
    ws.cell(row, 1).font = Font(name="微软雅黑", size=8, color="888888")

# 年度: B173-B176 (10Y, 15Y, 20Y, 30Y)
for i, (rate_row, t) in enumerate([(108, 10), (109, 15), (110, 20), (111, 30)]):
    row = 173 + i
    formula = f"=1/(1+$B${rate_row})^{t}"
    ws.cell(row, 2, formula)
    ws.cell(row, 2).font = Font(name="微软雅黑", size=9, color="555555")
    ws.cell(row, 2).number_format = "0.000000000"
    ws.cell(row, 2).border = border
    ws.cell(row, 1, f"DF {t}Y")
    ws.cell(row, 1).font = Font(name="微软雅黑", size=8, color="888888")

# 标题
ws.cell(112, 2, "折现因子辅助区")
ws.cell(112, 2).font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
ws.cell(112, 2).fill = PatternFill("solid", fgColor="C7000B")
ws.cell(112, 2).alignment = Alignment(horizontal="center", vertical="center")

# ===================== 3. 注释 =====================
# R111 下方添加公式说明
ws.cell(178, 2, "公式说明: B113-B172 = 1/(1+B48-B107)^(m/12) 月度折现因子")
ws.cell(178, 2).font = Font(name="微软雅黑", size=9, italic=True, color="0070C0")
ws.cell(179, 2, "          B173-B176 = 1/(1+B108-B111)^(10/15/20/30) 年度折现因子")
ws.cell(179, 2).font = Font(name="微软雅黑", size=9, italic=True, color="0070C0")
ws.cell(180, 2, "提示: WPS 打开后如折现因子为 0, 请按 Ctrl+Alt+F9 强制重算")
ws.cell(180, 2).font = Font(name="微软雅黑", size=9, bold=True, color="C7000B")

# ===================== 4. 保存 =====================
wb.save(EXCEL_PATH)
print(f"已保存: {EXCEL_PATH}")
print("B113-B176 改回公式 (v3 方式)")
print("K 列保持数值 (v4 方式, 立即可见)")

# 验证
print("\n=== 验证 B113-B116 公式 ===")
wb2 = load_workbook(EXCEL_PATH, data_only=False)
ws2 = wb2['基础数据']
for r in range(113, 117):
    print(f"  B{r}: {ws2.cell(r, 2).value}")

print("\n=== K 列数值 (R7, R29) ===")
wb3 = load_workbook(EXCEL_PATH, data_only=True)
ws3 = wb3['基础数据']
for r in [7, 29]:
    name = (ws3.cell(r, 4).value or '').strip()[:22]
    pv = ws3.cell(r, 11).value
    print(f"  R{r} {name:<22} PV = {pv/1e8:>14.2f} 亿元")
