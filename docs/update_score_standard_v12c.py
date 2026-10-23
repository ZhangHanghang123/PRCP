"""
更新 评分标准_v12b.xlsx → v12c.xlsx
在「0-500分评分标准」sheet 的「评语」列后插入两列：
  - 区间下限分数  (score_lo)
  - 区间上限分数  (score_hi)
用于区间内线性插值。

实现要点（不用 openpyxl.insert_cols — 它会丢失合并单元格右边的数据）：
  1. 把整表的值 + 样式快照到内存
  2. 在内存中重排列序（插入两列）
  3. 重建合并单元格范围（列号偏移）
  4. 一次性写回

插值规则（按值排序，与「上一段得分」对齐）：
  越高越好 (LCR/NSFR/CET1/ROE/LR):
      区间下限分数 = 上一段得分（值更小）
      区间上限分数 = 本段得分
  越低越好 (EVE/IBL):
      区间下限分数 = 本段得分（值更小 = 分数更高）
      区间上限分数 = 上一段得分（值更大 = 分数更低）
  红线段（首段）: 上下限分相同 = 本段得分
  单点段 (min == max): 上下限分相同 = 本段得分
"""
import shutil
from pathlib import Path
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = Path(r"C:/银行经营/PRCP/docs/v2/评分标准_v12b.xlsx")
DST = Path(r"C:/银行经营/PRCP/docs/v2/评分标准_v12c.xlsx")
SHEET = "0-500分评分标准"

TITLE_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="C7000B")
NEW_FILL = PatternFill("solid", fgColor="FFF3E0")
BOLD = Font(name="微软雅黑", size=10, bold=True)

THIN = Side(style="thin", color="DDDDDD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def parse_pct(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip().replace('%', ''))
    except ValueError:
        return None


def snapshot(ws):
    """整表值 + 样式快照"""
    snap = []
    for r in range(1, ws.max_row + 1):
        row = []
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            row.append({
                "value": cell.value,
                "font": copy(cell.font),
                "fill": copy(cell.fill),
                "border": copy(cell.border),
                "align": copy(cell.alignment),
                "numfmt": cell.number_format,
            })
        snap.append(row)
    return snap


def col_to_idx(s):
    n = 0
    for ch in s:
        n = n * 26 + (ord(ch.upper()) - ord('A') + 1)
    return n


def main():
    shutil.copy(SRC, DST)
    wb = load_workbook(DST)
    ws = wb[SHEET]

    old_ncol, old_nrow = ws.max_column, ws.max_row
    old_merges = [str(m) for m in ws.merged_cells.ranges]
    old_freeze = ws.freeze_panes
    old_widths = {k: v.width for k, v in ws.column_dimensions.items()}

    # ---------- 定位表头 + 得分列 ----------
    hdr_row = next(r for r in range(1, 10) if ws.cell(r, 1).value == "指标")
    score_col = next(c for c in range(1, old_ncol + 1)
                     if ws.cell(hdr_row, c).value == "得分")
    print(f"表头行=R{hdr_row}  得分列={get_column_letter(score_col)}({score_col})  原列数={old_ncol}")

    snap = snapshot(ws)

    # ---------- 1. 收集数据，按指标分组 ----------
    groups, cur = [], None
    for r in range(hdr_row + 1, old_nrow + 1):
        seg = ws.cell(r, 3).value
        lo, hi = parse_pct(ws.cell(r, 4).value), parse_pct(ws.cell(r, 5).value)
        sc = ws.cell(r, score_col).value
        if seg is None or lo is None or hi is None or sc is None:
            continue

        metric = ws.cell(r, 1).value
        if metric:
            direction = ws.cell(r, 10).value or ""
            cur = {"metric": metric, "direction": direction,
                   "higher": ("越高越好" in direction), "rows": []}
            groups.append(cur)
        if cur is None:
            continue
        cur["rows"].append({"row": r, "seg": int(seg), "lo": lo, "hi": hi,
                            "score": float(sc)})

    print(f"\n识别 {len(groups)} 个指标:")
    for g in groups:
        print(f"  {g['metric']:<6} {len(g['rows']):>2} 段  {g['direction']}")

    # ---------- 2. 计算上下限分数 ----------
    for g in groups:
        ordered = sorted(g["rows"], key=lambda x: x["lo"], reverse=not g["higher"])
        prev = None
        for i, it in enumerate(ordered):
            if it["lo"] == it["hi"] or i == 0:
                it["score_lo"] = it["score_hi"] = it["score"]
            elif g["higher"]:
                it["score_lo"], it["score_hi"] = prev, it["score"]
            else:
                it["score_lo"], it["score_hi"] = it["score"], prev
            prev = it["score"]

    # ---------- 3. 重排列序（在「评语」列后插 2 列） ----------
    NEW_LO = score_col + 2
    NEW_HI = score_col + 3
    new_ncol = old_ncol + 2

    new_snap = []
    for row in snap:
        nr = []
        for c in range(1, old_ncol + 1):
            if c == NEW_LO:
                nr.append(None)
                nr.append(None)
            nr.append(row[c - 1])
        new_snap.append(nr)

    # ---------- 4. 写回 ----------
    ws.delete_cols(1, ws.max_column)
    for r in range(1, old_nrow + 1):
        for c in range(1, new_ncol + 1):
            item = new_snap[r - 1][c - 1]
            cell = ws.cell(r, c)
            if item is None:
                continue
            cell.value = item["value"]
            cell.font = item["font"]
            cell.fill = item["fill"]
            cell.border = item["border"]
            cell.alignment = item["align"]
            cell.number_format = item["numfmt"]

    # 新增两列表头
    for c, title in ((NEW_LO, "区间下限分数"), (NEW_HI, "区间上限分数")):
        cell = ws.cell(hdr_row, c)
        cell.value = title
        cell.font = TITLE_FONT
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    # 新增两列数据
    missing = 0
    for g in groups:
        for it in g["rows"]:
            r = it["row"]
            for c, v in ((NEW_LO, it["score_lo"]), (NEW_HI, it["score_hi"])):
                cell = ws.cell(r, c)
                cell.value = v
                cell.font = BOLD
                cell.fill = NEW_FILL
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.number_format = "0"
                cell.border = BORDER

    # ---------- 5. 重建合并单元格（列号偏移） ----------
    ws.merged_cells.ranges = []

    def shift(ref):
        col_s = "".join(ch for ch in ref if ch.isalpha())
        row_s = "".join(ch for ch in ref if ch.isdigit())
        ci = col_to_idx(col_s)
        if ci >= NEW_LO:
            ci += 2
        return f"{get_column_letter(ci)}{row_s}"

    for m in old_merges:
        parts = m.split(":")
        new_refs = [shift(p) for p in parts]
        try:
            ws.merge_cells(":".join(new_refs))
        except Exception as e:
            print(f"  ⚠️ 跳过合并 {m} → {':'.join(new_refs)}: {e}")

    # ---------- 6. 列宽 / 冻结 ----------
    for col, w in old_widths.items():
        ci = col_to_idx(col)
        if ci >= NEW_LO:
            ci += 2
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.column_dimensions[get_column_letter(NEW_LO)].width = 13
    ws.column_dimensions[get_column_letter(NEW_HI)].width = 13
    ws.freeze_panes = old_freeze or "A5"

    wb.save(DST)

    # ---------- 7. 验证 ----------
    wb2 = load_workbook(DST, data_only=True)
    ws2 = wb2[SHEET]
    hdrs = [ws2.cell(hdr_row, c).value for c in range(1, ws2.max_column + 1)]
    print(f"\n{'='*78}")
    print(f"新表头 ({ws2.max_column} 列): {hdrs}")
    print("=" * 78)

    ci = {h: i + 1 for i, h in enumerate(hdrs) if h}
    c_lo, c_hi = ci["区间下限分数"], ci["区间上限分数"]
    c_seg, c_desc = ci["子段"], ci["档位"]

    print(f"\n列位置: 得分={get_column_letter(ci['得分'])}  评语={get_column_letter(ci['得分']+1)}  "
          f"区间下限分数={get_column_letter(c_lo)}  区间上限分数={get_column_letter(c_hi)}")

    for g in groups:
        print(f"\n{g['metric']} ({g['direction']})")
        print(f"  {'段':<5}{'区间':<24}{'下限分':<10}{'上限分':<10}档位")
        for it in g["rows"]:
            if it["seg"] not in (1, 2, 3, 6, 11, 12, 15, 16, 17, 18, 22, 23):
                continue
            r = it["row"]
            slo, shi = ws2.cell(r, c_lo).value, ws2.cell(r, c_hi).value
            tier = ws2.cell(r, c_desc).value or ""
            bad = "" if (slo is not None and shi is not None) else "  ❌缺失"
            if bad:
                missing += 1
            print(f"  {it['seg']:<5}[{it['lo']:g}%, {it['hi']:g}%]{'':<8}"
                  f"{slo if slo is not None else '?':<10}{shi if shi is not None else '?':<10}{tier}{bad}")

    # 评语列完整性检查
    desc_col = ci["得分"] + 1
    lost = sum(1 for g in groups for it in g["rows"]
               if not ws2.cell(it["row"], desc_col).value)
    print(f"\n数据完整性: 上下限分缺失 {missing} 处, 评语丢失 {lost} 处")

    print(f"\n✅ 已保存: {DST}")


if __name__ == "__main__":
    main()